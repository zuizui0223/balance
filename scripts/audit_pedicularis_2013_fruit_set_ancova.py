"""Fixed Table 1 fruit-set ANCOVA replay; partial matches remain a HOLD.

Run from the repository root with python -m scripts.audit_pedicularis_2013_fruit_set_ancova.
The optional pedicularis extra provides read-only XLSX extraction with openpyxl.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from scripts.adjudicate_pedicularis_2013_raw_workbook import (
    SOURCE_BYTES, SOURCE_MD5, SOURCE_SHA256, FRUIT_SHEET, SEED_SHEET,
    EXPECTED_FRUIT_HEADER, EXPECTED_COUNTS_FRUIT, _source_rows,
    ancova_fruit_set, compare_fruit_set_ancova,
)


def load_fruit_rows(workbook_file: Path):
    """Hash-verified, read-only 74-row source loader shared by audits."""
    content = workbook_file.read_bytes()
    if (len(content) != SOURCE_BYTES
            or hashlib.sha256(content).hexdigest() != SOURCE_SHA256
            or hashlib.md5(content).hexdigest() != SOURCE_MD5):
        raise ValueError("original Dryad source workbook bytes/digest changed")
    from openpyxl import load_workbook
    workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["pollination rate", FRUIT_SHEET, SEED_SHEET]:
            raise ValueError("original workbook sheet order changed")
        cells = [list(row) for row in workbook[FRUIT_SHEET].iter_rows(
            min_row=1, max_row=77, min_col=1, max_col=6, values_only=True)]
    finally:
        workbook.close()
    return _source_rows(cells[0], cells[1:], EXPECTED_FRUIT_HEADER,
                        EXPECTED_COUNTS_FRUIT)


def audit_source(workbook_file: Path) -> dict:
    rows, counts = load_fruit_rows(workbook_file)
    model = ancova_fruit_set(rows)
    if model["n"] != 74 or model["df_residual"] != 69:
        raise ValueError("pooled fruit-set ANCOVA source support changed")
    return {
        "schema_version": "PEDICULARIS_2013_FRUIT_SET_ANCOVA_AUDIT_V1",
        **compare_fruit_set_ancova(model["F"]),
        "source_article_doi": "10.1098/rsbl.2013.0387",
        "source_dataset_doi": "10.5061/dryad.6cv06",
        "source_workbook_sha256": SOURCE_SHA256,
        "source_workbook_md5": SOURCE_MD5,
        "source_workbook_bytes": SOURCE_BYTES,
        "original_source_bytes_verified": True,
        "source_sheet": FRUIT_SHEET,
        "source_biological_excel_rows": [2, 75],
        "source_group_counts": counts,
        "model": model,
        "model_formula": "recorded fruit-set percent ~ 1 + year + density + year:density + centered recorded flowers",
        "factor_coding": "year and density use 2*source_code-3; interaction is their product",
        "conditional_covariate_assumption": "recorded flowers represents the published plant-size covariate",
        "original_SPSS_specification_independently_verified": False,
        "analysis_status": "post-publication computational audit; no pre-outcome ecological confirmation",
        "number_of_specifications_tested": 1,
        "fruit_production_counts_reconstructed": False,
        "observed_rates_changed_or_imputed": False,
        "fruit_predation_ANCOVA_reproduced": False,
        "patch_clustered_inference_identified": False,
        "causal_density_effect_identified": False,
        "source_paper_error_established": False,
        "existing_24F_receipt_scope_changed": False,
        "claim_ceiling": "fixed conditional source-statistic replay only; retain every mismatch; no original-code or patch-level causal identification",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    if args.out.exists():
        raise ValueError("refusing to overwrite fruit-set ANCOVA source audit")
    result = audit_source(args.workbook)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("x", encoding="utf-8") as handle:
        handle.write(json.dumps(result, sort_keys=True, indent=2) + "\n")
    print(json.dumps({key: result[key] for key in
                     ("status", "matched_F_count", "attempted_F_count",
                      "all_F_reproduced")}, sort_keys=True))


if __name__ == "__main__":
    main()
