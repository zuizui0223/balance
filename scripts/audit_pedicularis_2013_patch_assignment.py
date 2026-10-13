"""Audit 2011 Pedicularis latent-patch average bounds from exact Dryad bytes.

No surrogate patch IDs are invented. The output is an OUTER mathematical
envelope over hypothetical assignments, NOT sampling uncertainty or selection.
"""
from __future__ import annotations
import argparse
import hashlib
import json
import tempfile
from pathlib import Path
from zipfile import ZipFile

from balance_domain.pedicularis_patch_assignment_bounds import density_contrast_envelope
from scripts.adjudicate_pedicularis_2013_raw_workbook import (
    SOURCE_BYTES, SOURCE_MD5, SOURCE_SHA256, SEED_SHEET, FRUIT_SHEET,
    EXPECTED_SEED_HEADER, EXPECTED_COUNTS_SEED, _source_rows,
)
from scripts.audit_pedicularis_2013_fruit_set_ancova import load_fruit_rows
from scripts.retrieve_pedicularis_2013_dryad_bytes import verify_zip

PUBLISHED_2011_PATCHES={"sparse":5,"dense":6}
EXPECTED_FRUIT_2011={"sparse":26,"dense":32}
EXPECTED_SEED_2011={"sparse":1070,"dense":1279}
OUTCOMES={
    "fruit_set":("fruit",4,[1,2,5]),
    "fruit_predation":("fruit",5,[1,2,5]),
    "final_seed_set":("seed",4,[1,5,20]),
    "seed_predation":("seed",5,[1,5,20]),
}


def load_seed_rows(workbook_file:Path):
    content=workbook_file.read_bytes()
    if (len(content)!=SOURCE_BYTES or hashlib.md5(content).hexdigest()!=SOURCE_MD5
            or hashlib.sha256(content).hexdigest()!=SOURCE_SHA256):
        raise ValueError("original 2013 Dryad source workbook bytes changed")
    from openpyxl import load_workbook
    w=load_workbook(workbook_file,read_only=True,data_only=True)
    try:
        if w.sheetnames!=["pollination rate",FRUIT_SHEET,SEED_SHEET]:
            raise ValueError("unexpected original sheet order")
        cells=[list(row) for row in w[SEED_SHEET].iter_rows(
            min_row=1,max_row=2933,min_col=1,max_col=6,values_only=True)]
    finally:
        w.close()
    rows,counts=_source_rows(cells[0],cells[1:],EXPECTED_SEED_HEADER,
                             EXPECTED_COUNTS_SEED)
    if sum(row[3] is None for row in rows)!=328:
        raise ValueError("initial-seed missingness pattern changed")
    return rows,counts


def audit_source(workbook_file:Path):
    fruit, fruit_counts=load_fruit_rows(workbook_file)
    seed, seed_counts=load_seed_rows(workbook_file)
    groups={}
    for source,rows,expected in (
        ("fruit",fruit,EXPECTED_FRUIT_2011),
        ("seed",seed,EXPECTED_SEED_2011),
    ):
        groups[source]={}
        for den,label in ((1,"sparse"),(2,"dense")):
            selected=[r for r in rows if r[0]==2 and r[1]==den]
            if len(selected)!=expected[label]:
                raise ValueError("2011 source density support changed")
            groups[source][label]=selected
    analyses={}
    for label,(source,icol,mins) in OUTCOMES.items():
        sample=groups[source]
        sparse=[r[icol] for r in sample["sparse"]]
        dense=[r[icol] for r in sample["dense"]]
        if any(v is None for v in sparse+dense):
            raise ValueError("missing final response cannot be imputed")
        analyses[label]={str(m): density_contrast_envelope(
            sparse,dense,PUBLISHED_2011_PATCHES["sparse"],
            PUBLISHED_2011_PATCHES["dense"],m) for m in mins}
    return {
        "schema_version":"PEDICULARIS_2013_RELAXED_LATENT_PATCH_MEAN_ENVELOPE_V1",
        "status":"SOURCE_VERIFIED_RELAXED_PATCH_ASSIGNMENT_OUTER_BOUND_NOT_EMPIRICAL_PATCH_INFERENCE",
        "dataset_doi":"10.5061/dryad.6cv06",
        "article_doi":"10.1098/rsbl.2013.0387",
        "source_workbook_sha256":SOURCE_SHA256,
        "source_workbook_bytes":SOURCE_BYTES,
        "original_source_bytes_verified":True,
        "analysis_year":2011,
        "published_patch_counts_by_density":PUBLISHED_2011_PATCHES,
        "original_2011_fruit_row_support":EXPECTED_FRUIT_2011,
        "original_2011_seed_row_support":EXPECTED_SEED_2011,
        "fruit_source_category_counts":fruit_counts,
        "seed_source_category_counts":seed_counts,
        "outcomes":analyses,
        "only_density_and_published_patch_counts_constrained":True,
        "actual_patch_ids_present":False,
        "plant_ids_present":False,
        "reported_patch_size_binary_labels_not_constrained_in_relaxed_envelope":True,
        "hypothetical_patch_assignments_are_not_actual_sample_design":True,
        "bounds_are_not_confidence_intervals_or_causal_effects":True,
        "source_supplemental_SPSS_code_recovered":False,
        "claim_ceiling":"unknown patch row assignments; extreme mathematical envelopes only; no cluster-robust p values, no BALANCE worldline inference",
    }


def audit_zip(path:Path):
    zipbytes=path.read_bytes()
    rawmeta=verify_zip(zipbytes)
    if rawmeta["source_sha256"]!=SOURCE_SHA256:
        raise ValueError("original biological workbook checksum changed")
    with ZipFile(path) as archive:
        content=archive.read("raw data.xlsx")
    with tempfile.TemporaryDirectory(prefix="pedicularis-patch-envelope-") as d:
        workbook=Path(d)/"raw data.xlsx"
        workbook.write_bytes(content)
        report=audit_source(workbook)
    report["downloaded_zip_sha256_observed_not_pinned"]=hashlib.sha256(zipbytes).hexdigest()
    return report


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--source-zip",type=Path,required=True)
    p.add_argument("--out",type=Path,required=True)
    a=p.parse_args()
    if a.out.exists():
        raise ValueError("refusing to overwrite patch envelope receipt")
    report=audit_zip(a.source_zip)
    a.out.parent.mkdir(parents=True,exist_ok=True)
    with a.out.open("x",encoding="utf-8") as f:
        json.dump(report,f,sort_keys=True,indent=2)
        f.write("\n")
    print(json.dumps(report,sort_keys=True))


if __name__=="__main__":
    main()
