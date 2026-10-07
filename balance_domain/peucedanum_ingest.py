"""Fail-closed source inventory and semantic mapping for Peucedanum raw data.

This layer deliberately separates byte/file inspection from biological semantics.
It never guesses source columns from their names. Normalized rows are emitted only
from an explicitly source-verified mapping contract.
"""
from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
import math
import os
import shutil
import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Iterable, Mapping

from .peucedanum_raw import REQUIRED_NORMALIZED_FIELDS, validate_normalized_rows


MAPPING_SCHEMA = "BALANCE_PEUCEDANUM_RAW_SEMANTIC_MAPPING_V1"
MAPPING_READY_STATUS = "SOURCE_VERIFIED_MAPPING"

INVENTORY_RECEIPT_BASENAME = "BALANCE_PEUCEDANUM_RAW_SOURCE_INVENTORY_V1.json"
NORMALIZED_ROWS_BASENAME = "BALANCE_PEUCEDANUM_NORMALIZED_ROWS_V1.csv"
NORMALIZATION_RECEIPT_BASENAME = "BALANCE_PEUCEDANUM_NORMALIZATION_RECEIPT_V1.json"

NORMALIZED_FIELDS = (
    "dataset_id",
    "source_doi",
    "source_file",
    "source_sheet",
    "source_row_number",
    "year",
    "population_id",
    "plant_id",
    "flowering_day",
    "perfect_flower_count",
    "male_flower_count",
    "initial_fruit_count",
    "intact_fruit_count",
    "predator_egg_count",
    "seed_predation_rate",
    "flower_stem_height",
    "male_fitness",
    "pollen_or_siring_measure",
    "total_flower_count",
    "male_fraction",
    "final_fruit_set_rate",
)

PROVENANCE_FIELDS = {
    "dataset_id",
    "source_doi",
    "source_file",
    "source_sheet",
    "source_row_number",
}

DERIVED_FIELDS = {
    "total_flower_count",
    "male_fraction",
    "final_fruit_set_rate",
}

TABULAR_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xls"}


def _sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _plain_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Peucedanum source contains a non-finite numeric cell")
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def _trim_matrix(rows: Iterable[Iterable[object]]) -> list[list[str]]:
    matrix = [[_plain_text(value) for value in row] for row in rows]
    while matrix and not any(matrix[-1]):
        matrix.pop()
    width = max((len(row) for row in matrix), default=0)
    matrix = [row + [""] * (width - len(row)) for row in matrix]
    while width and matrix and all(row[width - 1] == "" for row in matrix):
        width -= 1
    return [row[:width] for row in matrix]


def _table_inventory(
    *,
    source_file: str,
    source_sheet: str,
    rows: Iterable[Iterable[object]],
) -> dict:
    matrix = _trim_matrix(rows)
    width = max((len(row) for row in matrix), default=0)
    return {
        "source_file": source_file,
        "source_sheet": source_sheet,
        "row_count": len(matrix),
        "column_count": width,
        "nonempty_row_count": sum(any(cell != "" for cell in row) for row in matrix),
    }


def inventory_tables(
    tables: Mapping[tuple[str, str], Iterable[Iterable[object]]],
) -> list[dict]:
    """Inventory decoded tables without interpreting any biological column."""
    if not tables:
        raise ValueError("at least one source table is required")
    out = []
    for (source_file, source_sheet), rows in sorted(tables.items()):
        if not source_file or not source_sheet:
            raise ValueError("source file and sheet names must be non-empty")
        out.append(
            _table_inventory(
                source_file=source_file,
                source_sheet=source_sheet,
                rows=rows,
            )
        )
    return out


def _read_delimited(raw: bytes, *, delimiter: str) -> list[list[str]]:
    text = raw.decode("utf-8-sig")
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(raw: bytes) -> dict[str, list[list[object]]]:
    try:
        import openpyxl  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "Peucedanum .xlsx ingest requires openpyxl; install with "
            "python -m pip install -e '.[peucedanum]'"
        ) from exc

    book = openpyxl.load_workbook(
        filename=io.BytesIO(raw),
        read_only=True,
        data_only=False,
    )
    try:
        return {
            sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
            for sheet in book.worksheets
        }
    finally:
        book.close()


def _read_xls(raw: bytes) -> dict[str, list[list[object]]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "Peucedanum .xls ingest requires xlrd; install with "
            "python -m pip install -e '.[peucedanum]'"
        ) from exc

    book = xlrd.open_workbook(file_contents=raw, on_demand=True)
    sheets: dict[str, list[list[object]]] = {}
    try:
        for name in book.sheet_names():
            sheet = book.sheet_by_name(name)
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        value: object = ""
                    elif cell.ctype == xlrd.XL_CELL_TEXT:
                        value = str(cell.value).strip()
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        value = float(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        dt = xlrd.xldate_as_datetime(cell.value, book.datemode)
                        value = (
                            dt.date().isoformat()
                            if dt.time().isoformat() == "00:00:00"
                            else dt.isoformat()
                        )
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        value = f"#XL_ERROR_{int(cell.value)}"
                    else:
                        value = cell.value
                    row.append(value)
                rows.append(row)
            sheets[name] = rows
    finally:
        book.release_resources()
    return sheets


def _read_tabular_bytes(
    source_name: str,
    raw: bytes,
) -> dict[str, list[list[object]]]:
    suffix = Path(source_name).suffix.lower()
    if suffix == ".csv":
        return {"__CSV__": _read_delimited(raw, delimiter=",")}
    if suffix == ".tsv":
        return {"__TSV__": _read_delimited(raw, delimiter="\t")}
    if suffix == ".xlsx":
        return _read_xlsx(raw)
    if suffix == ".xls":
        return _read_xls(raw)
    raise ValueError(
        "Peucedanum tabular ingest accepts only .csv, .tsv, .xlsx or .xls files"
    )


def read_tabular_source(
    path: Path,
) -> tuple[bytes, dict[str, list[list[object]]]]:
    """Read one standalone tabular source without assigning biological semantics."""
    raw = path.read_bytes()
    if path.suffix.lower() not in TABULAR_SUFFIXES:
        raise ValueError(
            "Peucedanum source ingest accepts only .csv, .tsv, .xlsx or .xls files"
        )
    return raw, _read_tabular_bytes(path.name, raw)


def _safe_zip_members(raw: bytes) -> list[tuple[str, bytes]]:
    members: list[tuple[str, bytes]] = []
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = [info.filename for info in archive.infolist() if not info.is_dir()]
        if len(names) != len(set(names)):
            raise ValueError("Peucedanum ZIP contains duplicate member names")
        for info in archive.infolist():
            if info.is_dir():
                continue
            member = PurePosixPath(info.filename)
            if (
                member.is_absolute()
                or ".." in member.parts
                or "\\" in info.filename
                or not member.name
            ):
                raise ValueError(
                    f"Peucedanum ZIP contains unsafe member path {info.filename!r}"
                )
            data = archive.read(info)
            members.append((member.as_posix(), data))
    return sorted(members, key=lambda item: item[0])


def load_sources(
    source_paths: Iterable[Path],
) -> tuple[dict[str, dict], dict[tuple[str, str], list[list[object]]]]:
    """Load exact source/archive bytes and return provenance plus decoded tables."""
    paths = list(source_paths)
    if not paths:
        raise ValueError("at least one Peucedanum source file is required")
    names = [path.name for path in paths]
    if len(names) != len(set(names)):
        raise ValueError("Peucedanum source basenames must be unique")

    files: dict[str, dict] = {}
    tables: dict[tuple[str, str], list[list[object]]] = {}
    for path in paths:
        raw = path.read_bytes()
        if path.suffix.lower() == ".zip":
            members = _safe_zip_members(raw)
            member_inventory = [
                {
                    "member": member,
                    "bytes": len(member_raw),
                    "sha256": _sha256_bytes(member_raw),
                    "format": Path(member).suffix.lower().lstrip("."),
                }
                for member, member_raw in members
            ]
            files[path.name] = {
                "source_bytes": len(raw),
                "source_sha256": _sha256_bytes(raw),
                "format": "zip",
                "member_count": len(members),
                "members": member_inventory,
            }
            for member, member_raw in members:
                if Path(member).suffix.lower() not in TABULAR_SUFFIXES:
                    continue
                source_key = f"{path.name}::{member}"
                decoded = _read_tabular_bytes(member, member_raw)
                for sheet, rows in decoded.items():
                    tables[(source_key, sheet)] = rows
            continue

        decoded = _read_tabular_bytes(path.name, raw)
        files[path.name] = {
            "source_bytes": len(raw),
            "source_sha256": _sha256_bytes(raw),
            "format": path.suffix.lower().lstrip("."),
            "sheet_names": list(decoded),
        }
        for sheet, rows in decoded.items():
            tables[(path.name, sheet)] = rows
    return files, tables


def build_source_inventory_receipt(
    source_paths: Iterable[Path],
) -> dict:
    files, tables = load_sources(source_paths)
    return {
        "schema_version": "BALANCE_PEUCEDANUM_RAW_SOURCE_INVENTORY_V1",
        "status": "SOURCE_BYTES_INVENTORIED_SEMANTIC_MAPPING_NOT_APPLIED",
        "files": files,
        "tables": inventory_tables(tables),
        "semantic_mapping_applied": False,
        "source_hash_binding_verified": False,
        "claim_ceiling": "source_inventory_only_no_biological_result",
    }


def _source_registry(mapping: Mapping[str, object]) -> dict[str, Mapping[str, object]]:
    if mapping.get("schema_version") != MAPPING_SCHEMA:
        raise ValueError("Peucedanum semantic mapping schema mismatch")
    if mapping.get("status") != MAPPING_READY_STATUS:
        raise ValueError(
            "Peucedanum semantic mapping must be SOURCE_VERIFIED_MAPPING before normalization"
        )
    sources = mapping.get("registered_sources")
    if not isinstance(sources, list) or not sources:
        raise ValueError("Peucedanum mapping requires registered_sources")

    out: dict[str, Mapping[str, object]] = {}
    for source in sources:
        if not isinstance(source, Mapping):
            raise ValueError("registered source must be an object")
        doi = str(source.get("source_doi") or "").strip()
        if not doi or doi in out:
            raise ValueError("registered source DOI must be non-empty and unique")
        out[doi] = source
    return out


def _mapping_entries(mapping: Mapping[str, object]) -> list[dict]:
    sources = _source_registry(mapping)
    entries = []
    for doi, source in sources.items():
        raw_entries = source.get("sheet_mappings")
        if not isinstance(raw_entries, list):
            raise ValueError(
                f"registered source {doi!r} sheet_mappings must be a list"
            )
        if not raw_entries:
            continue
        expected = source.get("expected_files")
        if expected is None and source.get("expected_file"):
            expected = [source["expected_file"]]
        expected_files = (
            {str(name).strip() for name in expected}
            if isinstance(expected, list)
            else None
        )

        for raw_entry in raw_entries:
            if not isinstance(raw_entry, Mapping):
                raise ValueError("sheet mapping must be an object")
            entry = dict(raw_entry)
            source_file = str(entry.get("source_file") or "").strip()
            if expected_files is not None:
                exact_or_member = (
                    source_file in expected_files
                    or any(
                        source_file.startswith(f"{expected}::")
                        for expected in expected_files
                    )
                )
                if not exact_or_member:
                    raise ValueError(
                        f"sheet mapping source_file {source_file!r} is not registered "
                        f"for source {doi!r}"
                    )
            entry["source_doi"] = doi
            entries.append(entry)
    if not entries:
        raise ValueError("Peucedanum mapping contains no sheet mappings")
    return entries


def _verify_mapping_source_hashes(
    mapping: Mapping[str, object],
    files: Mapping[str, Mapping[str, object]],
) -> None:
    """Bind a source-verified semantic mapping to the exact outer source bytes."""
    sources = _source_registry(mapping)
    for doi, source in sources.items():
        raw_entries = source.get("sheet_mappings")
        if not isinstance(raw_entries, list):
            raise ValueError(
                f"registered source {doi!r} sheet_mappings must be a list"
            )
        if not raw_entries:
            continue

        raw_hashes = source.get("expected_file_sha256s")
        if not isinstance(raw_hashes, Mapping) or not raw_hashes:
            raise ValueError(
                f"registered source {doi!r} requires expected_file_sha256s "
                "before production normalization"
            )

        expected_hashes: dict[str, str] = {}
        for filename, digest in raw_hashes.items():
            name = str(filename).strip()
            value = str(digest).strip().lower()
            if (
                not name
                or len(value) != 64
                or any(ch not in "0123456789abcdef" for ch in value)
            ):
                raise ValueError(
                    f"registered source {doi!r} has invalid expected SHA256 binding"
                )
            expected_hashes[name] = value

        referenced_outer_files: set[str] = set()
        for raw_entry in raw_entries:
            if not isinstance(raw_entry, Mapping):
                raise ValueError("sheet mapping must be an object")
            source_file = str(raw_entry.get("source_file") or "").strip()
            if not source_file:
                raise ValueError("sheet mapping requires source_file")
            referenced_outer_files.add(source_file.split("::", 1)[0])

        missing = sorted(referenced_outer_files - set(expected_hashes))
        if missing:
            raise ValueError(
                f"registered source {doi!r} lacks expected SHA256 for mapped file(s): "
                + ", ".join(missing)
            )

        for outer in sorted(referenced_outer_files):
            info = files.get(outer)
            if info is None:
                raise ValueError(
                    f"source-verified mapping expects source file {outer!r} "
                    "but those bytes were not loaded"
                )
            actual = str(info.get("source_sha256") or "").strip().lower()
            if actual != expected_hashes[outer]:
                raise ValueError(
                    f"source SHA256 mismatch for {outer!r}: "
                    f"expected {expected_hashes[outer]}, got {actual}"
                )


def _header_index(header: list[str]) -> dict[str, int]:
    names = [name.strip() for name in header]
    if not names or any(not name for name in names):
        raise ValueError("mapped source header must not contain blank columns")
    if len(names) != len(set(names)):
        raise ValueError("mapped source header columns must be unique")
    return {name: i for i, name in enumerate(names)}


def _derive_registered_fields(row: dict[str, str]) -> None:
    def number(field: str) -> Decimal | None:
        value = row.get(field, "")
        if value == "":
            return None
        try:
            out = Decimal(value)
        except (InvalidOperation, ValueError) as exc:
            raise ValueError(f"mapped {field} is not numeric") from exc
        if not out.is_finite():
            raise ValueError(f"mapped {field} is not finite")
        return out

    perfect = number("perfect_flower_count")
    male = number("male_flower_count")
    if perfect is not None and male is not None:
        total = perfect + male
        if not total.is_finite() or total <= 0:
            raise ValueError("mapped flower counts produce invalid total")
        row["total_flower_count"] = format(total, "g")
        row["male_fraction"] = format(male / total, ".15g")

    intact = number("intact_fruit_count")
    if intact is not None and perfect is not None and perfect > 0:
        row["final_fruit_set_rate"] = format(intact / perfect, ".15g")


def normalize_from_mapping(
    tables: Mapping[tuple[str, str], Iterable[Iterable[object]]],
    mapping: Mapping[str, object],
) -> list[dict[str, str]]:
    """Apply only explicit source-verified mappings, then validate normalized rows."""
    entries = _mapping_entries(mapping)
    normalized: list[dict[str, str]] = []

    for entry in entries:
        dataset_id = str(entry.get("dataset_id") or "").strip()
        source_doi = str(entry["source_doi"]).strip()
        source_file = str(entry.get("source_file") or "").strip()
        source_sheet = str(entry.get("source_sheet") or "").strip()
        header_row = entry.get("header_row")
        columns = entry.get("columns")
        constants = entry.get("constants", {})

        if not dataset_id:
            raise ValueError("sheet mapping requires dataset_id")
        if not source_file or not source_sheet:
            raise ValueError("sheet mapping requires source_file and source_sheet")
        if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
            raise ValueError("sheet mapping header_row must be a positive integer")
        if not isinstance(columns, Mapping) or not columns:
            raise ValueError("sheet mapping requires explicit columns")
        if not isinstance(constants, Mapping):
            raise ValueError("sheet mapping constants must be an object")

        key = (source_file, source_sheet)
        if key not in tables:
            raise ValueError(f"mapped source table not loaded: {key!r}")
        matrix = _trim_matrix(tables[key])
        if header_row > len(matrix):
            raise ValueError(f"mapped header row exceeds source table: {key!r}")

        header = matrix[header_row - 1]
        index = _header_index(header)

        field_sources: dict[str, int] = {}
        for normalized_field, source_column in columns.items():
            normalized_field = str(normalized_field)
            source_column = str(source_column)
            if normalized_field not in NORMALIZED_FIELDS:
                raise ValueError(
                    f"unsupported normalized field {normalized_field!r}"
                )
            if normalized_field in PROVENANCE_FIELDS | DERIVED_FIELDS:
                raise ValueError(
                    f"field {normalized_field!r} is provenance/derived and cannot be mapped"
                )
            if source_column not in index:
                raise ValueError(
                    f"mapped source column {source_column!r} not found in {key!r}"
                )
            field_sources[normalized_field] = index[source_column]

        constant_values = {
            str(field): _plain_text(value)
            for field, value in constants.items()
        }

        raw_missing_tokens = entry.get("missing_tokens", [])
        if not isinstance(raw_missing_tokens, list):
            raise ValueError("sheet mapping missing_tokens must be a list")
        missing_tokens = {_plain_text(value) for value in raw_missing_tokens}

        raw_value_maps = entry.get("value_maps", {})
        if not isinstance(raw_value_maps, Mapping):
            raise ValueError("sheet mapping value_maps must be an object")
        value_maps: dict[str, dict[str, str]] = {}
        for field, raw_map in raw_value_maps.items():
            field = str(field)
            if field not in field_sources:
                raise ValueError(
                    f"value map field {field!r} must also have a mapped source column"
                )
            if not isinstance(raw_map, Mapping) or not raw_map:
                raise ValueError(
                    f"value map for {field!r} must be a non-empty object"
                )
            mapped: dict[str, str] = {}
            for source_value, normalized_value in raw_map.items():
                source_text = _plain_text(source_value)
                normalized_text = _plain_text(normalized_value)
                if source_text in mapped:
                    raise ValueError(
                        f"value map for {field!r} repeats source value {source_text!r}"
                    )
                mapped[source_text] = normalized_text
            value_maps[field] = mapped
        if any(field not in NORMALIZED_FIELDS for field in constant_values):
            raise ValueError("sheet mapping contains unsupported constant field")
        if any(field in PROVENANCE_FIELDS for field in constant_values):
            raise ValueError("provenance normalized fields cannot be supplied as constants")
        if any(field in DERIVED_FIELDS for field in constant_values):
            raise ValueError("derived normalized fields cannot be supplied as constants")
        overlap = set(field_sources) & set(constant_values)
        if overlap:
            raise ValueError(
                "normalized fields cannot be supplied by both column and constant: "
                + ", ".join(sorted(overlap))
            )

        provided = set(field_sources) | set(constant_values) | {
            "dataset_id",
            "source_doi",
        }
        missing_required = [
            field for field in REQUIRED_NORMALIZED_FIELDS if field not in provided
        ]
        if missing_required:
            raise ValueError(
                f"mapping {dataset_id!r} lacks required normalized fields: "
                + ", ".join(missing_required)
            )

        for row_index, raw_row in enumerate(matrix[header_row:], start=header_row + 1):
            if not any(cell != "" for cell in raw_row):
                continue
            padded = raw_row + [""] * (len(header) - len(raw_row))
            if len(padded) > len(header):
                raise ValueError(f"source row {row_index} exceeds mapped header width")

            row = {field: "" for field in NORMALIZED_FIELDS}
            row.update({
                "dataset_id": dataset_id,
                "source_doi": source_doi,
                "source_file": source_file,
                "source_sheet": source_sheet,
                "source_row_number": str(row_index),
            })
            row.update(constant_values)
            for field, col_index in field_sources.items():
                source_value = padded[col_index]
                if source_value in missing_tokens:
                    row[field] = ""
                    continue
                if field in value_maps:
                    if source_value not in value_maps[field]:
                        raise ValueError(
                            f"mapped source value {source_value!r} for {field!r} "
                            "is not registered in value_maps"
                        )
                    row[field] = value_maps[field][source_value]
                else:
                    row[field] = source_value
            _derive_registered_fields(row)
            normalized.append(row)

    if not normalized:
        raise ValueError("source-verified mapping produced no normalized rows")
    validate_normalized_rows(normalized)
    return normalized


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=NORMALIZED_FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def inventory_peucedanum_sources(
    source_paths: Iterable[Path],
    out_dir: Path,
) -> dict:
    """Write an immutable R0 byte/table inventory without semantic mapping."""
    if out_dir.exists():
        raise ValueError("Peucedanum inventory output directory already exists")
    receipt = build_source_inventory_receipt(source_paths)
    out_dir.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix=f".{out_dir.name}.tmp-", dir=out_dir.parent))
    try:
        (tmp / INVENTORY_RECEIPT_BASENAME).write_text(
            json.dumps(receipt, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        os.replace(tmp, out_dir)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise
    return receipt


def normalize_peucedanum_sources(
    source_paths: Iterable[Path],
    mapping_path: Path,
    out_dir: Path,
) -> dict:
    """Normalize source rows only from a reviewed mapping contract."""
    if out_dir.exists():
        raise ValueError("Peucedanum normalization output directory already exists")
    source_paths = list(source_paths)
    files, tables = load_sources(source_paths)
    mapping_raw = mapping_path.read_bytes()
    mapping = json.loads(mapping_raw.decode("utf-8"))
    _verify_mapping_source_hashes(mapping, files)
    rows = normalize_from_mapping(tables, mapping)
    inventory = inventory_tables(tables)

    receipt = {
        "schema_version": "BALANCE_PEUCEDANUM_NORMALIZATION_RECEIPT_V1",
        "status": "SOURCE_VERIFIED_MAPPING_APPLIED",
        "mapping_file": mapping_path.name,
        "mapping_sha256": _sha256_bytes(mapping_raw),
        "source_hash_binding_verified": True,
        "source_files": files,
        "source_tables": inventory,
        "normalized_rows": len(rows),
        "normalized_output": NORMALIZED_ROWS_BASENAME,
        "next_gate": "reproduce_published_2025_source_models_before_new_criticality",
        "claim_ceiling": "normalized_raw_rows_only_no_criticality_result",
    }

    out_dir.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix=f".{out_dir.name}.tmp-", dir=out_dir.parent))
    try:
        _write_csv(tmp / NORMALIZED_ROWS_BASENAME, rows)
        (tmp / NORMALIZATION_RECEIPT_BASENAME).write_text(
            json.dumps(receipt, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        os.replace(tmp, out_dir)
    except Exception:
        shutil.rmtree(tmp, ignore_errors=True)
        raise
    return receipt
